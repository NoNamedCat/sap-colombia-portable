"""Pruebas TDD para el Dashboard Gerencial, paginación, filtros y exportación de reportes."""

import json
from sap_colombia import dbutil
import threading
import time
from decimal import Decimal
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest
from http.server import HTTPServer

from sap_colombia import database as db
from sap_colombia import auth
from sap_colombia import excel_export
from sap_colombia import inventory
from sap_colombia.web_server import SapWebHandler

TEST_PORT = 8091


@pytest.fixture(autouse=True)
def setup_db():
    db.init_db()
    conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            "TRUNCATE TABLE movimientos_contables, recibos_caja, factura_lineas, "
            "lineas_nota_credito, notas_credito, asientos_contables, facturas, "
            "kardex, stock_actual CASCADE"
        )
        conn.commit()
    finally:
        conn.close()
    yield


@pytest.fixture(scope="module")
def web_server():
    db.init_db()
    auth.seed_usuarios()
    from sap_colombia import rbac
    rbac.ensure_rbac()
    server = HTTPServer(("127.0.0.1", TEST_PORT), SapWebHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    time.sleep(0.5)
    yield f"http://127.0.0.1:{TEST_PORT}"
    server.shutdown()
    server.server_close()


def _admin_token() -> str:
    import json
    req = Request(
        f"http://127.0.0.1:{TEST_PORT}/api/auth/login",
        data=json.dumps({"email": "admin@saluki.com", "password": "admin123"}).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    with urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))["token"]


def _get(url: str) -> tuple[int, dict]:
    token = _admin_token()
    req = Request(url, headers={"Authorization": f"Bearer {token}"})
    with urlopen(req) as resp:
        return resp.status, json.loads(resp.read().decode("utf-8"))


def _get_raw(url: str) -> tuple[int, bytes]:
    token = _admin_token()
    req = Request(url, headers={"Authorization": f"Bearer {token}"})
    with urlopen(req) as resp:
        return resp.status, resp.read()


def _get_raw_allow_error(url: str) -> tuple[int, bytes]:
    token = _admin_token()
    req = Request(url, headers={"Authorization": f"Bearer {token}"})
    try:
        with urlopen(req) as resp:
            return resp.status, resp.read()
    except HTTPError as e:
        return e.code, e.read()


def _crear_factura(consecutivo: str, total: str, saldo: str, fecha: str, vencimiento: str, cliente: str = "800197268") -> int:
    return db.guardar_factura(
        {
            "consecutivo": consecutivo,
            "prefijo": "SETP",
            "fecha": fecha,
            "cliente_nit": cliente,
            "total_factura": total,
            "saldo_pendiente": saldo,
            "fecha_vencimiento": vencimiento,
            "total_iva": "0",
            "retefuente": "0",
        },
        [],
    )


# ===============================
# Tests de exportar_csv / exportar_excel
# ===============================

def test_exportar_csv_valido():
    rows = [
        {"id": 1, "cliente_nit": "800197268", "saldo": "1000.00"},
        {"id": 2, "cliente_nit": "900123456", "saldo": "2000.00"},
    ]
    data = excel_export.exportar_csv(["id", "cliente_nit", "saldo"], rows, ["ID", "Cliente", "Saldo"])
    texto = data.decode("utf-8-sig")
    lineas = texto.strip().splitlines()
    assert lineas[0] == "ID,Cliente,Saldo"
    assert lineas[1] == "1,800197268,1000.00"
    assert lineas[2] == "2,900123456,2000.00"


def test_exportar_csv_vacio():
    data = excel_export.exportar_csv(["id", "total"], [], ["ID", "Total"])
    texto = data.decode("utf-8-sig")
    assert texto.strip() == "ID,Total"


def test_exportar_excel_valido():
    if not excel_export.OPENPYXL_AVAILABLE:
        pytest.skip("openpyxl no disponible")
    rows = [{"id": 1, "total": "100.00"}]
    data = excel_export.exportar_excel(["id", "total"], rows, "Reporte")
    assert len(data) > 0
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.title == "Reporte"
    assert ws["A1"].value == "id"
    assert ws["B1"].value == "total"
    assert ws["A2"].value == 1


# ===============================
# Tests de paginación y filtros en /api/facturas
# ===============================

def test_facturas_paginacion(web_server):
    for i in range(1, 11):
        _crear_factura(str(i), "1000.00", "1000.00", f"2026-01-{i:02d}", f"2026-02-{i:02d}", f"800000000{i}")

    status, data = _get(f"{web_server}/api/facturas?page=1&limit=5")
    assert status == 200
    assert data["total"] == 10
    assert data["page"] == 1
    assert data["limit"] == 5
    assert data["pages"] == 2
    assert len(data["data"]) == 5


def test_facturas_filtros_fecha(web_server):
    for i in range(1, 4):
        _crear_factura(f"F{i}", "1000.00", "1000.00", "2026-01-10", "2026-02-10", f"9{i:08d}")

    status, data = _get(f"{web_server}/api/facturas?fecha_desde=2026-01-01&fecha_hasta=2026-01-31")
    assert status == 200
    assert data["total"] == 3

    status, data = _get(f"{web_server}/api/facturas?fecha_desde=2026-02-01&fecha_hasta=2026-02-28")
    assert status == 200
    assert data["total"] == 0


def test_facturas_filtros_cliente(web_server):
    _crear_factura("C1", "1000.00", "1000.00", "2026-01-01", "2026-02-01", "111111")
    _crear_factura("C2", "1000.00", "1000.00", "2026-01-01", "2026-02-01", "222222")

    status, data = _get(f"{web_server}/api/facturas?cliente_nit=111111")
    assert status == 200
    assert data["total"] == 1
    assert data["data"][0]["cliente_nit"] == "111111"


def test_facturas_filtros_estado_cartera(web_server):
    _crear_factura("E1", "1000.00", "1000.00", "2026-01-01", "2026-02-01", "333333")

    status, data = _get(f"{web_server}/api/facturas?estado_cartera=PENDIENTE")
    assert status == 200
    assert data["total"] == 1

    status, data = _get(f"{web_server}/api/facturas?estado_cartera=Liquidada")
    assert status == 200
    assert data["total"] == 0


# ===============================
# Tests de /api/dashboard/kpis
# ===============================

def test_dashboard_kpis_base(web_server):
    _crear_factura("K1", "500000.00", "300000.00", "2026-10-01", "2026-11-01", "444444")
    _crear_factura("K2", "200000.00", "200000.00", "2026-10-02", "2026-11-02", "555555")

    status, data = _get(f"{web_server}/api/dashboard/kpis")
    assert status == 200
    assert "ventas_mes" in data
    assert "cartera_pendiente" in data
    assert "cartera_vencida" in data
    assert "ticket_promedio" in data
    assert "stock" in data
    assert "contabilidad" in data

    assert data["ventas_mes"] == "700000.00"
    assert data["cartera_pendiente"] == "500000.00"
    assert data["ticket_promedio"] == "350000.00"


def test_dashboard_kpis_stock_alertas(web_server):
    conn = db.get_connection()
    try:
        emp_id = dbutil.query(conn, "SELECT id FROM empresas ORDER BY id LIMIT 1").fetchone()["id"]
        dbutil.query(conn, 
            """INSERT INTO productos (empresa_id, codigo, descripcion, precio_unitario, tarifa_iva, tipo_bien_servicio, stock_minimo)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT DO NOTHING""",
            (emp_id, "P-ALERTA", "Producto crítico", "100.00", "IVA_19", "BIENES", 10),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        raise
    finally:
        conn.close()

    inventory.registrar_entrada("P-ALERTA", "BODEGA-01", Decimal("5"), Decimal("50.00"))

    status, data = _get(f"{web_server}/api/dashboard/kpis")
    assert status == 200
    stock_alertas = data["stock"]["stock_alertas"]
    assert any(item["producto_codigo"] == "P-ALERTA" for item in stock_alertas)


# ===============================
# Tests de endpoints de exportación
# ===============================

def test_exportar_cartera_csv(web_server):
    _crear_factura("X1", "1000.00", "1000.00", "2026-10-01", "2026-11-01", "666666")

    status, data = _get_raw(f"{web_server}/api/reportes/cartera/exportar?formato=csv")
    assert status == 200
    texto = data.decode("utf-8-sig")
    assert "Cliente" in texto
    assert "800197268" not in texto  # cliente 800197268 no tiene saldo pendiente
    assert "666666" in texto


def test_exportar_diario_csv(web_server):
    status, data = _get_raw(f"{web_server}/api/reportes/libro-diario/exportar?formato=csv")
    assert status == 200
    texto = data.decode("utf-8-sig")
    assert "Fecha" in texto


def test_exportar_kardex_csv(web_server):
    inventory.registrar_entrada("P-KARD", "BODEGA-01", Decimal("10"), Decimal("100.00"))

    status, data = _get_raw(f"{web_server}/api/reportes/kardex/exportar?formato=csv")
    assert status == 200
    texto = data.decode("utf-8-sig")
    assert "Producto" in texto
    assert "P-KARD" in texto


def test_exportar_formato_invalido(web_server):
    status, data = _get_raw_allow_error(f"{web_server}/api/reportes/cartera/exportar?formato=pdf")
    assert status == 400


# ===============================
# Tests de paginación en otros endpoints
# ===============================

def test_cartera_edades_con_total(web_server):
    _crear_factura("CT1", "1000.00", "1000.00", "2026-10-01", "2026-11-01", "777777")

    status, data = _get(f"{web_server}/api/cartera/edades")
    assert status == 200
    assert data["total"] == 1
    assert len(data["detalle"]) == 1


def test_contabilidad_asientos_paginado(web_server):
    status, data = _get(f"{web_server}/api/contabilidad/asientos?page=1&limit=10")
    assert status == 200
    assert "total" in data
    assert "data" in data
